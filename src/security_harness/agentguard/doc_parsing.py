"""Local, bounded document extraction. Document text is always untrusted data."""

from __future__ import annotations

import base64
import binascii
import hashlib
import json
import zipfile
from dataclasses import dataclass, field
from io import BytesIO
from xml.etree import ElementTree as ET

from .detection_v2 import MAX_SCAN_CHARS, MAX_SEGMENTS, detect_indirect_segments

PdfReader = None
Document = None
_PARSERS_LOADED = False

MAX_DOCUMENT_BYTES = 8 * 1024 * 1024
MAX_XML_BYTES = 4 * 1024 * 1024
MAX_DOCUMENT_SEGMENTS = 4096
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@dataclass(frozen=True)
class DocumentSegment:
    doc_type: str
    part: str
    field: str
    text: str = field(repr=False)
    page: int | None = None

    @property
    def labels(self) -> dict:
        return {
            "doc_type": self.doc_type,
            "part": self.part,
            "page": self.page,
            "field": self.field,
        }


@dataclass(frozen=True)
class ParsedDocument:
    doc_type: str
    segments: tuple[DocumentSegment, ...] = ()
    status: str = "ok"
    code: str | None = None


def document_type(data: bytes, *, mime: str = "", filename: str = "") -> str | None:
    """Prefer magic; a generic ZIP alone does not establish DOCX."""
    if data.startswith(b"%PDF-"):
        return "pdf"
    if data.startswith(b"PK"):
        try:
            with zipfile.ZipFile(BytesIO(data)) as archive:
                if "word/document.xml" in archive.namelist():
                    return "docx"
        except zipfile.BadZipFile:
            pass
    mime = mime.lower().split(";", 1)[0].strip()
    filename = filename.lower().split("?", 1)[0].split("#", 1)[0]
    if mime == "application/pdf" or filename.endswith(".pdf"):
        return "pdf"
    if mime == DOCX_MIME or filename.endswith(".docx"):
        return "docx"
    return None


def parse_document(data: bytes, *, mime: str = "", filename: str = "") -> ParsedDocument:
    """Return explicit unavailable/error states; never present partial extraction as clean."""
    global PdfReader, Document, _PARSERS_LOADED
    if not _PARSERS_LOADED:
        try:
            from pypdf import PdfReader as _PdfReader
            PdfReader = _PdfReader
        except ImportError:
            pass
        try:
            from docx import Document as _Document
            Document = _Document
        except ImportError:
            pass
        _PARSERS_LOADED = True
    kind = document_type(data, mime=mime, filename=filename)
    if kind is None:
        return ParsedDocument("unknown", status="unsupported", code="DOCUMENT_UNSUPPORTED")
    if len(data) > MAX_DOCUMENT_BYTES:
        return ParsedDocument(kind, status="error", code="DOCUMENT_SIZE_LIMIT")
    if kind == "pdf" and PdfReader is None:
        return ParsedDocument(kind, status="unavailable", code="PDF_PARSER_UNAVAILABLE")
    segments = []
    total = 0

    def add(part, name, text, page=None):
        nonlocal total
        if text is None or not str(text).strip():
            return
        text = str(text)
        total += len(text)
        if total > MAX_SCAN_CHARS or len(segments) >= MAX_DOCUMENT_SEGMENTS:
            raise ValueError("document limit")
        segments.append(DocumentSegment(kind, part, name, text, page))

    try:
        if kind == "pdf":
            reader = PdfReader(BytesIO(data))
            for name in ("Title", "Author", "Subject", "Keywords", "Creator", "Producer"):
                add("metadata", name, (reader.metadata or {}).get("/" + name))
            # Scan the complete packet, including custom namespaces and attributes.
            xmp = reader.trailer["/Root"].get("/Metadata")
            if xmp is not None:
                add("metadata", "XMP", xmp.get_object().get_data().decode("utf-8"))
            for number, page in enumerate(reader.pages, 1):
                add("page", "text", page.extract_text(), number)
                annotations = page.get("/Annots", [])
                if hasattr(annotations, "get_object"):
                    annotations = annotations.get_object()
                for index, annotation in enumerate(annotations):
                    add(
                        "annotation",
                        f"contents[{index}]",
                        annotation.get_object().get("/Contents"),
                        number,
                    )
        else:
            with zipfile.ZipFile(BytesIO(data)) as archive:
                members = archive.infolist()
                if (
                    len(members) > MAX_DOCUMENT_SEGMENTS
                    or sum(m.file_size for m in members) > MAX_XML_BYTES
                    or len({m.filename for m in members}) != len(members)
                ):
                    raise ValueError("archive limit")
                for member in members:
                    name = member.filename
                    if not (
                        name == "docProps/core.xml"
                        or name == "word/document.xml"
                        or name.startswith(("word/header", "word/footer", "word/comments"))
                        and name.endswith(".xml")
                    ):
                        continue
                    xml = archive.read(member)
                    declarations = xml.upper().replace(b"\x00", b"")
                    if b"<!DOCTYPE" in declarations or b"<!ENTITY" in declarations:
                        raise ValueError("XML entities forbidden")
                    root = ET.fromstring(xml)
                    if name == "word/document.xml" and Document is not None:
                        # python-docx supplies the body tree; XML traversal also sees hidden runs.
                        root = Document(BytesIO(data)).element
                    if name == "docProps/core.xml":
                        for child in root:
                            add(name, child.tag.rsplit("}", 1)[-1], "".join(child.itertext()))
                        continue
                    table_paragraphs = {p for t in root.iter(W + "tbl") for p in t.iter(W + "p")}
                    for index, paragraph in enumerate(root.iter(W + "p")):
                        text = "".join(t.text or "" for t in paragraph.iter(W + "t"))
                        category = "table" if paragraph in table_paragraphs else "paragraph"
                        add(name, f"{category}[{index}]", text)
                        for run_index, run in enumerate(paragraph.iter(W + "r")):
                            vanish = run.find(f"{W}rPr/{W}vanish")
                            if vanish is not None and vanish.get(W + "val", "true") not in (
                                "0",
                                "false",
                                "off",
                            ):
                                add(
                                    name,
                                    f"hidden_text[{index}:{run_index}]",
                                    "".join(t.text or "" for t in run.iter(W + "t")),
                                )
    except Exception:
        # Parser exception strings can contain attacker-controlled document data.
        return ParsedDocument(kind, status="error", code="DOCUMENT_PARSE_FAILED")
    return ParsedDocument(kind, tuple(segments))


def scan_document(document: ParsedDocument, artifact_ids: list[str] | None = None) -> list:
    """Multi-source scans preserve labels through artifact IDs and conservative spans."""
    if document.status != "ok":
        raise ValueError(document.code)

    def identifier(value):
        return "0" + hashlib.sha256(value.encode()).hexdigest()[:25].upper()

    ids = (
        artifact_ids
        if artifact_ids is not None
        else [
            identifier(f"document:{index}:" + json.dumps(segment.labels, sort_keys=True))
            for index, segment in enumerate(document.segments)
        ]
    )
    if len(ids) != len(document.segments):
        raise ValueError("DOCUMENT_SOURCE_MISMATCH")
    segments = [(sid, segment.text) for sid, segment in zip(ids, document.segments)]
    findings = []
    # Also combine batch boundaries, without exceeding the detector's source count.
    # Each batch overlaps the preceding batch by one segment.
    for start in range(0, len(segments), MAX_SEGMENTS - 1):
        _, finding = detect_indirect_segments(
            segments[start : start + MAX_SEGMENTS], identifier(f"document-finding:{start}")
        )
        if finding is not None:
            findings.append(finding)
    return findings


def tool_documents(value: dict) -> list[tuple[tuple, ParsedDocument]]:
    """Find inline documents in JSON/MCP blocks; do not fetch filenames or URLs."""
    found = []

    def visit(item, path=(), inherited_mime="", inherited_name=""):
        if isinstance(item, dict):
            mime = item.get("mimeType", item.get("mime_type", item.get("mime", inherited_mime)))
            name = next(
                (
                    item[k]
                    for k in ("filename", "name", "resource", "path", "url", "uri")
                    if isinstance(item.get(k), str)
                ),
                inherited_name,
            )
            for key, child in item.items():
                if isinstance(child, (dict, list)):
                    visit(child, (*path, key), mime, name)
                elif key in {"content", "data", "blob", "bytes", "base64", "content_base64"}:
                    raw = child
                    content_mime = str(mime)
                    if isinstance(raw, str):
                        if raw.startswith("data:") and ";base64," in raw:
                            header, raw = raw.split(",", 1)
                            content_mime = header[5:].split(";", 1)[0]
                        try:
                            raw = base64.b64decode("".join(raw.split()), validate=True)
                        except (ValueError, binascii.Error):
                            raw = child.encode("utf-8")
                    if not isinstance(raw, bytes):
                        continue
                    if document_type(raw, mime=content_mime, filename=name):
                        found.append(
                            ((*path, key), parse_document(raw, mime=content_mime, filename=name))
                        )
        elif isinstance(item, list):
            for index, child in enumerate(item):
                visit(child, (*path, index), inherited_mime, inherited_name)

    visit(value)
    return found
